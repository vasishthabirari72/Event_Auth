import base64
import csv
import io
import zipfile
from typing import Any
from xml.etree.ElementTree import ParseError

from defusedxml.common import DefusedXmlException  # type: ignore[import-untyped]
from openpyxl import load_workbook  # type: ignore[import-untyped]

from event_auth.core.members.rules import RuleViolation


def read_rows(content: str, format: str) -> list[dict[str, str]]:
    try:
        raw = base64.b64decode(content, validate=True)
        if len(raw) > 2_000_000:
            raise ValueError
        rows: list[list[Any]] = []
        if format == "csv":
            for row in csv.reader(io.StringIO(raw.decode("utf-8-sig"))):
                rows.append(row)
                if len(rows) > 1001:
                    raise ValueError
        elif format == "xlsx":
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 20_000_000:
                    raise ValueError
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = book.worksheets[0]
                if (sheet.max_column or 0) > 30 or (sheet.max_row or 0) > 1001:
                    raise ValueError
                for cells in sheet.iter_rows():
                    if any(cell.data_type == "f" for cell in cells):
                        raise ValueError
                    rows.append([cell.value for cell in cells])
                    if len(rows) > 1001:
                        raise ValueError
            finally:
                book.close()
        else:
            raise ValueError
        if len(rows) < 2 or len(rows[0]) > 30:
            raise ValueError
        headers = [str(cell or "").strip() for cell in rows[0]]
        if not all(headers) or len(set(headers)) != len(headers):
            raise ValueError
        result = []
        for row in rows[1:]:
            if len(row) != len(headers):
                raise ValueError
            values = [str(cell if cell is not None else "").strip() for cell in row]
            if any(len(value) > 200 or value.startswith(("=", "+", "@")) for value in values):
                raise ValueError
            result.append(dict(zip(headers, values, strict=True)))
        return result
    except (
        ValueError,
        UnicodeError,
        zipfile.BadZipFile,
        KeyError,
        IndexError,
        OSError,
        DefusedXmlException,
        ParseError,
    ) as exc:
        raise RuleViolation("invalid_import") from exc
